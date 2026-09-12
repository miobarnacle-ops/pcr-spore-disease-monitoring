from pathlib import Path
import re
import sys
from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph


def iter_blocks(document):
    parent = document.element.body
    for child in parent.iterchildren():
        if child.tag == qn('w:p'):
            yield Paragraph(child, document)
        elif child.tag == qn('w:tbl'):
            yield Table(child, document)


def safe_name(name):
    stem = re.sub(r'[^0-9A-Za-z\u4e00-\u9fff._-]+', '_', name)
    return stem.strip('._') or 'image'


def extract(source, output_root):
    document = Document(source)
    target = output_root / source.stem
    image_dir = target / 'images'
    image_dir.mkdir(parents=True, exist_ok=True)
    lines = [f'# {source.name}', '']
    image_index = 0

    for block in iter_blocks(document):
        if isinstance(block, Paragraph):
            style = block.style.name if block.style else ''
            text = block.text.strip()
            level = 0
            match = re.search(r'(\d+)$', style)
            if style.lower().startswith('heading') and match:
                level = max(1, min(6, int(match.group(1))))
            if text:
                lines.append(('#' * level + ' ' if level else '') + text)
                lines.append('')
            for blip in block._p.xpath('.//a:blip'):
                rel_id = blip.get(qn('r:embed'))
                if not rel_id or rel_id not in document.part.rels:
                    continue
                part = document.part.rels[rel_id].target_part
                image_index += 1
                suffix = Path(str(part.partname)).suffix or '.bin'
                image_name = safe_name(f'{image_index:02d}_{Path(str(part.partname)).stem}{suffix}')
                image_path = image_dir / image_name
                image_path.write_bytes(part.blob)
                lines.append(f'![成员稿内嵌图片 {image_index}](images/{image_name})')
                lines.append('')
        else:
            for row in block.rows:
                values = [cell.text.replace('\n', ' / ').strip() for cell in row.cells]
                lines.append('| ' + ' | '.join(values) + ' |')
            lines.append('')

    metadata = [
        f'- 段落数：{len(document.paragraphs)}',
        f'- 表格数：{len(document.tables)}',
        f'- 图片数：{image_index}',
    ]
    lines[1:1] = metadata + ['']
    target.mkdir(parents=True, exist_ok=True)
    (target / 'content.md').write_text('\n'.join(lines), encoding='utf-8')
    print(f'{source.name}\tparagraphs={len(document.paragraphs)}\ttables={len(document.tables)}\timages={image_index}')


if __name__ == '__main__':
    output_root = Path(sys.argv[1]).resolve()
    for item in sys.argv[2:]:
        extract(Path(item).resolve(), output_root)
