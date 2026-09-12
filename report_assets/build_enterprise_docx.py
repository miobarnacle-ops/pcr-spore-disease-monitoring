from pathlib import Path
import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path(sys.argv[1]).resolve()
OUTPUT = Path(sys.argv[2]).resolve()
EQ_DIR = ROOT / "report_assets" / "equations"
FONT = "Noto Sans CJK SC"
NAVY = "17365D"
PALE = "F4F8FB"
BORDER = "D9D9D9"


def set_font(run, size=None, bold=None, italic=None, color="000000"):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    node = OxmlElement("w:shd")
    node.set(qn("w:fill"), fill)
    props.append(node)


def cell_borders(cell):
    props = cell._tc.get_or_add_tcPr()
    borders = props.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        props.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), BORDER)
        borders.append(node)


def cell_margins(cell, value=105):
    props = cell._tc.get_or_add_tcPr()
    margins = OxmlElement("w:tcMar")
    for side in ("top", "start", "bottom", "end"):
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        margins.append(node)
    props.append(margins)


def repeat_header(row):
    props = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    props.append(node)


def set_table_geometry(table, widths_twips):
    props = table._tbl.tblPr
    width = props.first_child_found_in("w:tblW")
    if width is None:
        width = OxmlElement("w:tblW")
        props.append(width)
    width.set(qn("w:type"), "dxa")
    width.set(qn("w:w"), str(sum(widths_twips)))

    indent = props.first_child_found_in("w:tblInd")
    if indent is None:
        indent = OxmlElement("w:tblInd")
        props.append(indent)
    indent.set(qn("w:type"), "dxa")
    indent.set(qn("w:w"), "105")

    layout = props.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        props.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for value in widths_twips:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(value))
        grid.append(col)


def set_cell_width(cell, width_twips):
    props = cell._tc.get_or_add_tcPr()
    width = props.first_child_found_in("w:tcW")
    if width is None:
        width = OxmlElement("w:tcW")
        props.append(width)
    width.set(qn("w:type"), "dxa")
    width.set(qn("w:w"), str(width_twips))


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    text = OxmlElement("w:instrText"); text.set(qn("xml:space"), "preserve"); text.text = " PAGE "
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, text, end])
    set_font(run, 9, color="666666")


def clean(text):
    text = re.sub(r"\*\*(.*?)\*\*", r"\1", text.strip())
    text = re.sub(r"`([^`]*)`", r"\1", text)
    return text


def configure(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    section.different_first_page_header_footer = True
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = header.add_run("自主巡检导航 多端控制与孢子扩散预测技术方案")
    set_font(run, 8.5, color="666666")
    add_page_number(section.footer.paragraphs[0])
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2"):
        style = doc.styles[name]
        style.font.name = FONT
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
        style.font.color.rgb = RGBColor(0, 0, 0)
    normal = doc.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.45
    normal.paragraph_format.space_after = Pt(7)
    for name, size in (("Heading 1", 16), ("Heading 2", 13)):
        style = doc.styles[name]
        style.font.size = Pt(size)
        style.font.bold = True
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(14 if name == "Heading 1" else 10)
        style.paragraph_format.space_after = Pt(7)
    doc.core_properties.title = "自主巡检导航 多端控制与孢子扩散预测技术方案"
    doc.core_properties.subject = "项目技术方案模块整合稿"
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""


def add_cover(doc, title):
    for _ in range(5):
        doc.add_paragraph()
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(22)
    run = p.add_run(title)
    set_font(run, 23, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("项目技术方案模块整合稿")
    set_font(run, 15)
    p.paragraph_format.space_after = Pt(120)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("2026 年 9 月")
    set_font(run, 11, color="555555")
    doc.add_page_break()


def add_heading(doc, text, level):
    p = doc.add_paragraph(style=f"Heading {level}")
    run = p.add_run(clean(text))
    set_font(run, 16 if level == 1 else 13, bold=True)
    return p


def add_body(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Inches(0.29)
    p.paragraph_format.line_spacing = 1.45
    p.paragraph_format.space_after = Pt(7)
    run = p.add_run(clean(text))
    set_font(run, 11)
    return p


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = False
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run(clean(text.strip("*")))
    set_font(run, 9, italic=True, color="555555")


def add_image(doc, target):
    path = (ROOT / target).resolve()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(2)
    if "Android" in path.name or "手机" in path.name:
        width = Inches(2.35)
    elif path.parent.name == "equations":
        width = Inches(5.25)
    else:
        width = Inches(5.9)
    inline = p.add_run().add_picture(str(path), width=width)
    descr = ("公式 " if path.parent.name == "equations" else "技术图表 ") + path.stem
    inline._inline.docPr.set("descr", descr)
    inline._inline.docPr.set("title", descr)


def split_row(line):
    return [item.strip() for item in line.strip().strip("|").split("|")]


def separator(row):
    return all(re.fullmatch(r":?-{3,}:?", item.replace(" ", "")) for item in row)


def add_table(doc, source_rows):
    rows = [split_row(line) for line in source_rows]
    rows = [row for row in rows if not separator(row)]
    cols = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    usable = 6.94
    if cols == 3:
        widths = [usable * 0.23, usable * 0.30, usable * 0.47]
    elif cols == 4:
        widths = [usable * 0.18, usable * 0.25, usable * 0.31, usable * 0.26]
    else:
        widths = [usable / cols] * cols
    widths_twips = [round(value * 1440) for value in widths]
    set_table_geometry(table, widths_twips)
    for r_index, values in enumerate(rows):
        for c_index in range(cols):
            cell = table.cell(r_index, c_index)
            set_cell_width(cell, widths_twips[c_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell_margins(cell)
            cell_borders(cell)
            shade(cell, NAVY if r_index == 0 else (PALE if r_index % 2 == 0 else "FFFFFF"))
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_index < 2 else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.2
            run = p.add_run(clean(values[c_index]) if c_index < len(values) else "")
            set_font(run, 9.2 if r_index == 0 else 9,
                     bold=(r_index == 0), color="FFFFFF" if r_index == 0 else "000000")
    repeat_header(table.rows[0])
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def equation_path(block):
    if "begin{bmatrix}" in block:
        return EQ_DIR / "field_to_map.png"
    if "c_q" in block or "C_S" in block:
        return EQ_DIR / "sampling_coverage.png"
    if "C(x,y,z)" in block:
        return EQ_DIR / "gaussian_plume.png"
    if "partial C" in block:
        return EQ_DIR / "dynamic_advection_diffusion.png"
    if "P_{inf}" in block:
        return EQ_DIR / "infection_probability.png"
    return None


def build():
    doc = Document()
    configure(doc)
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    pending = []
    cover_done = False
    i = 0

    def flush():
        nonlocal pending
        if pending:
            add_body(doc, " ".join(part.strip() for part in pending))
            pending = []

    while i < len(lines):
        stripped = lines[i].strip()
        if not stripped:
            flush(); i += 1; continue
        if stripped.startswith("!["):
            flush()
            match = re.search(r"\]\(([^)]+)\)", stripped)
            if match:
                add_image(doc, match.group(1))
            i += 1; continue
        if stripped.startswith("\\["):
            flush(); block = [stripped]
            while i + 1 < len(lines) and not lines[i].strip().endswith("\\]"):
                i += 1; block.append(lines[i].strip())
            path = equation_path(" ".join(block))
            if path and path.exists():
                add_image(doc, path.relative_to(ROOT))
            i += 1; continue
        if stripped.startswith("|"):
            flush(); rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(lines[i].strip()); i += 1
            add_table(doc, rows); continue
        if stripped.startswith("#"):
            flush()
            marks = len(stripped) - len(stripped.lstrip("#"))
            heading = stripped[marks:].strip()
            if not cover_done:
                add_cover(doc, heading); cover_done = True
            elif marks == 1:
                doc.add_page_break(); add_heading(doc, heading, 1)
            else:
                add_heading(doc, heading, 2)
            i += 1; continue
        if re.match(r"^\*(?:图|表).+\*$", stripped):
            flush(); add_caption(doc, stripped); i += 1; continue
        pending.append(stripped); i += 1
    flush()
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
