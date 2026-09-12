from pathlib import Path
import re
import sys
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent.parent
SOURCE = Path(sys.argv[1]).resolve() if len(sys.argv) > 1 else ROOT / '报告书三模块初稿_图表数据与多端成果更新版_v4_2026-09-11.md'
OUTPUT = Path(sys.argv[2]).resolve() if len(sys.argv) > 2 else ROOT / '三模块技术方案初稿_图表数据与多端成果更新版_2026-09-11.docx'
APPEND_SOURCE = Path(sys.argv[3]).resolve() if len(sys.argv) > 3 else None
EQ_DIR = ROOT / 'report_assets' / 'equations'

FONT = 'Noto Sans CJK SC'
BLUE = '1F4E78'
LIGHT_BLUE = 'D9EAF7'
VERY_LIGHT = 'F5F9FC'
GRAY = '666666'


def set_run_font(run, size=None, bold=None, italic=None, color=None):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tc_pr.append(shd)


def set_cell_border(cell, color='B7C9D6'):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tc_pr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right'):
        tag = 'w:' + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), '4')
        element.set(qn('w:space'), '0')
        element.set(qn('w:color'), color)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in('w:tcMar')
    if tc_mar is None:
        tc_mar = OxmlElement('w:tcMar')
        tc_pr.append(tc_mar)
    for m, value in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tc_mar.find(qn('w:' + m))
        if node is None:
            node = OxmlElement('w:' + m)
            tc_mar.append(node)
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement('w:tblHeader')
    tbl_header.set(qn('w:val'), 'true')
    tr_pr.append(tbl_header)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run('第 ')
    set_run_font(run, 9, color=(90, 90, 90))
    fld_char1 = OxmlElement('w:fldChar')
    fld_char1.set(qn('w:fldCharType'), 'begin')
    instr_text = OxmlElement('w:instrText')
    instr_text.set(qn('xml:space'), 'preserve')
    instr_text.text = 'PAGE'
    fld_char2 = OxmlElement('w:fldChar')
    fld_char2.set(qn('w:fldCharType'), 'end')
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    run2 = paragraph.add_run(' 页')
    set_run_font(run2, 9, color=(90, 90, 90))


def clean_inline(text):
    text = text.strip()
    text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)
    text = re.sub(r'`([^`]*)`', r'\1', text)
    return text


def add_text_paragraph(doc, text, kind='body'):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.32
    if kind == 'quote':
        p.paragraph_format.left_indent = Inches(0.18)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(clean_inline(text))
        set_run_font(r, 9.5, italic=True, color=(88, 88, 88))
    elif kind == 'caption':
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(9)
        r = p.add_run(clean_inline(text.strip('*')))
        set_run_font(r, 9, italic=True, color=(80, 80, 80))
    elif kind == 'bullet':
        p.paragraph_format.left_indent = Inches(0.26)
        p.paragraph_format.first_line_indent = Inches(-0.16)
        r = p.add_run('• ' + clean_inline(text))
        set_run_font(r, 10.5)
    else:
        r = p.add_run(clean_inline(text))
        set_run_font(r, 10.5)
    return p


def add_heading(doc, text, level):
    p = doc.add_paragraph()
    p.style = f'Heading {level}'
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(15 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(clean_inline(text))
    if level == 1:
        set_run_font(r, 16, bold=True, color=(0, 0, 0))
    elif level == 2:
        set_run_font(r, 13, bold=True, color=(0, 0, 0))
    else:
        set_run_font(r, 11.5, bold=True, color=(0, 0, 0))
    return p


def add_image(doc, image_path):
    path = Path(image_path)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    name = path.name
    if 'Android' in name or 'android' in name or '手机' in name:
        width = Inches(2.55)
    elif '截图' in name:
        width = Inches(5.35)
    elif path.parent.name == 'equations':
        width = Inches(5.5)
    else:
        width = Inches(6.05)
    inline = p.add_run().add_picture(str(path), width=width)
    description = '公式图：' + path.stem if path.parent.name == 'equations' else '技术图表：' + path.stem
    inline._inline.docPr.set('descr', description)
    inline._inline.docPr.set('title', description)
    return p


def split_table_row(line):
    return [part.strip() for part in line.strip().strip('|').split('|')]


def is_separator(row):
    return all(re.fullmatch(r':?-{3,}:?', cell.replace(' ', '')) for cell in row)


def add_table(doc, rows):
    parsed = [split_table_row(line) for line in rows]
    parsed = [row for row in parsed if not is_separator(row)]
    if not parsed:
        return
    columns = max(len(row) for row in parsed)
    table = doc.add_table(rows=len(parsed), cols=columns)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    for row_i, row_data in enumerate(parsed):
        for col_i in range(columns):
            cell = table.cell(row_i, col_i)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            set_cell_border(cell)
            if row_i == 0:
                shade(cell, BLUE)
            elif row_i % 2 == 0:
                shade(cell, VERY_LIGHT)
            else:
                shade(cell, 'FFFFFF')
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            content = clean_inline(row_data[col_i]) if col_i < len(row_data) else ''
            r = p.add_run(content)
            if row_i == 0:
                set_run_font(r, 9.2, bold=True, color=(255, 255, 255))
            else:
                set_run_font(r, 8.7)
    set_repeat_table_header(table.rows[0])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def configure_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.66)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.68)
    section.right_margin = Inches(0.68)
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = header.add_run('基于机器人与 DNA-PCR 技术的孢子病害识别关键技术研究')
    set_run_font(r, 8.5, color=(90, 90, 90))
    add_page_number(section.footer.paragraphs[0])
    for style_name in ['Normal', 'Title', 'Heading 1', 'Heading 2', 'Heading 3']:
        style = doc.styles[style_name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
        style.font.color.rgb = RGBColor(0, 0, 0)
    title_ppr = doc.styles['Title']._element.get_or_add_pPr()
    title_border = title_ppr.find(qn('w:pBdr'))
    if title_border is not None:
        title_ppr.remove(title_border)


def build():
    doc = Document()
    configure_document(doc)
    lines = SOURCE.read_text(encoding='utf-8').splitlines()
    if APPEND_SOURCE is not None:
        append_lines = APPEND_SOURCE.read_text(encoding='utf-8').splitlines()
        section_two = next((idx for idx, value in enumerate(append_lines) if value.startswith('# 二、')), None)
        if section_two is None:
            raise ValueError('追加源文件中未找到第二章标题')
        lines.extend(append_lines[section_two:])
    first_title = True
    pending = []
    i = 0

    def flush_pending():
        nonlocal pending
        if pending:
            text = ' '.join(part.strip() for part in pending).strip()
            if text:
                add_text_paragraph(doc, text)
            pending = []

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if not stripped:
            flush_pending()
            i += 1
            continue
        if stripped == '---':
            flush_pending()
            i += 1
            continue
        if stripped.startswith('!['):
            flush_pending()
            match = re.search(r'\]\(([^)]+)\)', stripped)
            if match:
                target = match.group(1)
                image_path = ROOT / target
                if image_path.exists():
                    add_image(doc, image_path)
            i += 1
            continue
        if stripped.startswith('\\['):
            flush_pending()
            block = [stripped]
            while i + 1 < len(lines) and not lines[i].strip().endswith('\\]'):
                i += 1
                block.append(lines[i].strip())
            equation_text = ' '.join(block)
            equation_path = None
            if 'begin{bmatrix}' in equation_text:
                equation_path = EQ_DIR / 'field_to_map.png'
            elif 'c_q' in equation_text or 'C_S' in equation_text:
                equation_path = EQ_DIR / 'sampling_coverage.png'
            elif 'C(x,y,z)' in equation_text:
                equation_path = EQ_DIR / 'gaussian_plume.png'
            elif 'partial C' in equation_text:
                equation_path = EQ_DIR / 'dynamic_advection_diffusion.png'
            elif 'P_{inf}' in equation_text:
                equation_path = EQ_DIR / 'infection_probability.png'
            if equation_path is not None and equation_path.exists():
                add_image(doc, equation_path)
            i += 1
            continue
        if stripped.startswith('|'):
            flush_pending()
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                table_lines.append(lines[i].strip())
                i += 1
            add_table(doc, table_lines)
            continue
        if stripped.startswith('#'):
            flush_pending()
            marks = len(stripped) - len(stripped.lstrip('#'))
            heading = stripped[marks:].strip()
            if marks == 1 and first_title:
                p = doc.add_paragraph()
                p.style = 'Title'
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(10)
                ppr = p._p.get_or_add_pPr()
                border = ppr.find(qn('w:pBdr'))
                if border is not None:
                    ppr.remove(border)
                r = p.add_run('三模块技术方案证据补强初稿')
                set_run_font(r, 21, bold=True, color=(0, 0, 0))
                first_title = False
            else:
                if marks == 1:
                    doc.add_page_break()
                add_heading(doc, heading, min(marks, 3))
            i += 1
            continue
        if stripped.startswith('>'):
            flush_pending()
            add_text_paragraph(doc, stripped.lstrip('>').strip(), 'quote')
            i += 1
            continue
        if re.match(r'^\*(?:图|表).*\*$', stripped):
            flush_pending()
            add_text_paragraph(doc, stripped, 'caption')
            i += 1
            continue
        if re.match(r'^(?:[-*]|\d+\.)\s+', stripped):
            flush_pending()
            text = re.sub(r'^(?:[-*]|\d+\.)\s+', '', stripped)
            add_text_paragraph(doc, text, 'bullet')
            i += 1
            continue
        pending.append(stripped)
        i += 1

    flush_pending()
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == '__main__':
    build()
